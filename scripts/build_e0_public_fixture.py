#!/usr/bin/env python3
"""Build the binary assets for the public E0 contract fixture.

This fixture-only generator is self-contained within the repository. It uses
the MIT-licensed python-docx default template, Pillow, lxml, and standard
OPC/WordprocessingML operations to create synthetic content and a returned
document with fixed review events. It intentionally does not import Codex
skills or any other private helper. Redistributed template provenance is
recorded in the fixture manifest and ``third_party/python-docx-LICENSE.txt``.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import shutil
import sys
import zipfile
from pathlib import Path

import docx
import PIL
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from PIL import Image, ImageDraw

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
XML_NS = "http://www.w3.org/XML/1998/namespace"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
CP_NS = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
DC_NS = "http://purl.org/dc/elements/1.1/"
EP_NS = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
NS = {
    "w": W_NS,
    "cp": CP_NS,
    "dc": DC_NS,
    "ep": EP_NS,
    "pr": PKG_REL_NS,
    "ct": CT_NS,
}

COMMENTS_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
COMMENTS_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"
)

REVIEWER_ALPHA = "Reviewer Alpha"
REVIEWER_BETA = "Reviewer Beta"
TIME_ALPHA = "2026-01-15T09:30:00+08:00"
TIME_BETA = "2026-01-15T10:05:00-05:00"

TARGETS = {
    "txr_insert_001": "Insertion case: The test confirms repeatable behavior.",
    "txr_delete_001": "Deletion case: The redundant modifier remains visible in the baseline.",
    "txr_replace_001": "Replacement case: The response was good during the cycle.",
    "txr_move_001": (
        "Move case: Sensors were calibrated before loading. The specimen was then loaded. "
    ),
    "txr_format_001": (
        "Formatting case: The phrase critical observation is intentionally plain in the baseline."
    ),
    "txr_comment_range_001": (
        "Range comment anchor: verify the phrase normalized residual before approval."
    ),
    "txr_comment_point_001": ("Point comment anchor: check the following sentence boundary."),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--project-root",
        type=Path,
        default=Path(__file__).resolve().parents[1],
        help="Repository root (default: inferred from this script)",
    )
    parser.add_argument(
        "--fixture-root",
        type=Path,
        default=None,
        help="Fixture directory (default: tests/fixtures/e0-minimal-paper)",
    )
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def assert_within(root: Path, candidate: Path, label: str) -> Path:
    resolved_root = root.resolve()
    resolved_candidate = candidate.resolve()
    try:
        resolved_candidate.relative_to(resolved_root)
    except ValueError as exc:
        raise ValueError(
            f"{label} must remain inside {resolved_root}: {resolved_candidate}"
        ) from exc
    return resolved_candidate


def qnw(local: str) -> str:
    return f"{{{W_NS}}}{local}"


def deterministic_zip_members(members: dict[str, bytes], destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(destination, "w") as archive:
        for name in sorted(members):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            archive.writestr(info, members[name])


def read_zip_members(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path, "r") as archive:
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise ValueError(f"duplicate ZIP member in {path}")
        return {name: archive.read(name) for name in names}


def parse_xml_bytes(data: bytes, part_name: str) -> etree._Element:
    declaration_scan = data.upper().replace(b"\x00", b"")
    if b"<!DOCTYPE" in declaration_scan or b"<!ENTITY" in declaration_scan:
        raise ValueError(f"DTD or entity declaration is forbidden in {part_name}")
    parser = etree.XMLParser(
        resolve_entities=False,
        load_dtd=False,
        no_network=True,
        recover=False,
        remove_blank_text=False,
        huge_tree=False,
    )
    return etree.fromstring(data, parser=parser)


def xml_bytes(root: etree._Element) -> bytes:
    return etree.tostring(
        root,
        xml_declaration=True,
        encoding="UTF-8",
        standalone="yes",
    )


def sanitize_public_docx(source: Path, destination: Path) -> None:
    """Remove template-only and identifying metadata from a generated DOCX.

    This is a bounded fixture-generation pass over a package we just created;
    it is not the production privacy or ingest implementation. All operations
    are ordinary OPC relationship/content-type edits and WordprocessingML
    attribute removal.
    """

    members = read_zip_members(source)
    names = set(members)
    removed = {
        name
        for name in names
        if name.startswith("customXml/")
        or name.lower() in {"docprops/thumbnail.jpeg", "docprops/custom.xml"}
    }

    for name in sorted(names - removed):
        if not name.endswith((".xml", ".rels")):
            continue
        root = parse_xml_bytes(members[name], name)
        changed = False

        for element in root.iter():
            for attribute_name in list(element.attrib):
                attribute = etree.QName(attribute_name)
                if attribute.namespace == W_NS and attribute.localname.lower().startswith("rsid"):
                    del element.attrib[attribute_name]
                    changed = True

        if name == "docProps/core.xml":
            for xpath in ("dc:creator", "cp:lastModifiedBy"):
                node = root.find(xpath, namespaces=NS)
                if node is not None and node.text:
                    node.text = None
                    changed = True
        elif name == "docProps/app.xml":
            for xpath in ("ep:Company", "ep:Manager", "ep:HyperlinkBase"):
                node = root.find(xpath, namespaces=NS)
                if node is not None and node.text:
                    node.text = None
                    changed = True

        if name.endswith(".rels"):
            for relationship in list(root):
                target = (relationship.get("Target") or "").replace("\\", "/").lower()
                rel_type = (relationship.get("Type") or "").lower()
                if (
                    "customxml/" in target
                    or "thumbnail" in target
                    or target.endswith("docprops/custom.xml")
                    or rel_type.endswith(("/customxml", "/custom-properties"))
                ):
                    root.remove(relationship)
                    changed = True

        if changed:
            members[name] = xml_bytes(root)

    content_types = "[Content_Types].xml"
    ct_root = parse_xml_bytes(members[content_types], content_types)
    for child in list(ct_root):
        part_name = (child.get("PartName") or "").lower()
        extension = (child.get("Extension") or "").lower()
        if (
            part_name.startswith("/customxml/")
            or part_name in {"/docprops/thumbnail.jpeg", "/docprops/custom.xml"}
            or extension in {"jpeg", "jpg"}
            and not any(
                item.lower().endswith((".jpeg", ".jpg")) and item not in removed for item in names
            )
        ):
            ct_root.remove(child)
    members[content_types] = xml_bytes(ct_root)

    for name in removed:
        members.pop(name, None)
    deterministic_zip_members(members, destination)


def generate_response_curve(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1200, 620), "white")
    draw = ImageDraw.Draw(image)
    axis = (40, 55, 75)
    grid = (220, 226, 232)
    curve = (46, 116, 181)
    accent = (155, 28, 28)

    left, top, right, bottom = 120, 70, 1120, 525
    for i in range(6):
        y = top + i * (bottom - top) // 5
        draw.line((left, y, right, y), fill=grid, width=2)
    for i in range(11):
        x = left + i * (right - left) // 10
        draw.line((x, top, x, bottom), fill=grid, width=2)
    draw.line((left, top, left, bottom), fill=axis, width=5)
    draw.line((left, bottom, right, bottom), fill=axis, width=5)

    points = []
    values = [0.08, 0.18, 0.31, 0.50, 0.69, 0.84, 0.93, 0.98, 1.00, 0.97, 0.91]
    for idx, value in enumerate(values):
        x = left + idx * (right - left) // 10
        y = bottom - int(value * (bottom - top - 25))
        points.append((x, y))
    draw.line(points, fill=curve, width=8, joint="curve")
    for x, y in points:
        draw.ellipse((x - 9, y - 9, x + 9, y + 9), fill=accent, outline="white", width=3)

    image.save(path, format="PNG", optimize=False, compress_level=9)


def set_style_font(style, name: str, size_pt: float, color: str = "000000", bold=None) -> None:
    style.font.name = name
    style.font.size = Pt(size_pt)
    style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def apply_design_preset(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    set_style_font(normal, "Calibri", 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    title = styles["Title"]
    set_style_font(title, "Calibri", 16, color="0B2545", bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = styles["Subtitle"]
    set_style_font(subtitle, "Calibri", 10, color="555555")
    subtitle.font.italic = True
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading_specs = {
        "Heading 1": (16, "2E74B5", 18, 10),
        "Heading 2": (13, "2E74B5", 12, 6),
        "Heading 3": (12, "1F4D78", 8, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        set_style_font(style, "Calibri", size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    set_style_font(caption, "Calibri", 10, color="555555")
    caption.font.italic = True
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    insert_at = 1 if paragraph._p.pPr is not None else 0
    paragraph._p.insert(insert_at, start)
    paragraph._p.append(end)


def add_page_number_footer(document: Document) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Synthetic public fixture  |  ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, cached, end):
        run._r.append(node)


def add_native_math_paragraph(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math = parse_xml(
        f'<m:oMathPara xmlns:m="{M_NS}" xmlns:w="{W_NS}">'
        "<m:oMath><m:r><m:t>R = (Fmax - Fmin) / Fmax</m:t></m:r></m:oMath>"
        "</m:oMathPara>"
    )
    paragraph._p.append(math)


def mark_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def apply_table_geometry(
    table,
    column_widths_dxa: list[int],
    *,
    table_width_dxa: int,
    indent_dxa: int,
    cell_margins_dxa: dict[str, int],
) -> None:
    """Apply fixed OOXML table widths and cell margins.

    Word stores these values in twentieths of a point (DXA/twips). The helper
    is intentionally small and fixture-specific; it mirrors no private API.
    """

    if len(column_widths_dxa) != len(table.columns):
        raise ValueError("column width count does not match the table")

    table_properties = table._tbl.tblPr
    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(table_width_dxa))

    for tag in ("w:jc", "w:tblLayout", "w:tblInd"):
        existing = table_properties.find(qn(tag))
        if existing is not None:
            table_properties.remove(existing)

    width_index = table_properties.index(table_width)
    alignment = OxmlElement("w:jc")
    alignment.set(qn("w:val"), "left")
    table_properties.insert(width_index + 1, alignment)

    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table_properties.insert(width_index + 2, layout)

    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(indent_dxa))
    table_properties.append(indent)

    grid_columns = table._tbl.tblGrid.findall(qn("w:gridCol"))
    if len(grid_columns) != len(column_widths_dxa):
        raise ValueError("table grid column count does not match requested widths")
    for grid_column, width in zip(grid_columns, column_widths_dxa, strict=True):
        grid_column.set(qn("w:w"), str(width))

    for row in table.rows:
        row._tr.get_or_add_trPr()
        for column_index, cell in enumerate(row.cells):
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.insert(0, cell_width)
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(column_widths_dxa[column_index]))

            existing_margins = cell_properties.find(qn("w:tcMar"))
            if existing_margins is not None:
                cell_properties.remove(existing_margins)
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "start", "end"):
                margin = OxmlElement(f"w:{side}")
                margin.set(qn("w:w"), str(cell_margins_dxa[side]))
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cell_properties.append(margins)


def add_metrics_table(document: Document) -> None:
    document.add_paragraph("Table 1. Synthetic metrics / 合成指标", style="Caption")
    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    data = [
        ("Metric / 指标", "Baseline", "Reviewed"),
        ("Peak response / 峰值", "1.00", "1.05"),
        ("Residual ratio / 残余比", "0.12", "0.10"),
    ]
    for row_index, values in enumerate(data):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F4F6F9")
                cell._tc.get_or_add_tcPr().append(shading)
    mark_repeat_table_header(table.rows[0])
    apply_table_geometry(
        table,
        [3600, 2880, 2880],
        table_width_dxa=9360,
        indent_dxa=120,
        cell_margins_dxa={"top": 80, "bottom": 80, "start": 120, "end": 120},
    )


def build_base_docx(path: Path, image_path: Path) -> None:
    document = Document()
    apply_design_preset(document)
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = "E0 synthetic public review fixture"
    fixed_time = dt.datetime(2026, 1, 15, 0, 0, 0, tzinfo=dt.UTC)
    document.core_properties.created = fixed_time
    document.core_properties.modified = fixed_time

    document.add_paragraph(
        "A Reproducible LaTeX-Word Review Fixture / 可复现审阅样例", style="Title"
    )
    document.add_paragraph("Open Fixture Authors  |  2026-01-15", style="Subtitle")

    document.add_heading("Abstract / 摘要", level=1)
    document.add_paragraph(
        "This entirely synthetic document supports deterministic conversion and review tests. "
        "本文是完全合成的公开测试材料，用于验证中英文内容、公式、图表、修订与批注。"
    )

    document.add_heading("1. Introduction / 引言", level=1)
    document.add_paragraph(
        "The baseline states that the specimen response is stable under cyclic loading. "
        "本文基线说明试件在循环荷载下响应稳定。"
    )
    document.add_paragraph(
        "The measured response uses the inline relation E = mc², while the display equation "
        "defines the normalized residual."
    )

    document.add_heading("2. Method / 方法", level=1)
    document.add_paragraph(
        "The following native equation is intentionally simple so contract tests can count it."
    )
    add_native_math_paragraph(document)

    document.add_heading("3. Results / 结果", level=1)
    figure = document.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = figure.add_run().add_picture(str(image_path), width=Inches(5.4))
    shape._inline.docPr.set("descr", "Synthetic response curve")
    document.add_paragraph("Figure 1. Synthetic response curve / 合成响应曲线", style="Caption")
    add_metrics_table(document)

    document.add_heading("4. Review operations / 审阅操作", level=1)
    for bookmark_id, (unit_id, text) in enumerate(TARGETS.items(), start=10):
        paragraph = document.add_paragraph(text)
        add_bookmark(paragraph, unit_id, bookmark_id)

    document.add_heading("5. References / 参考文献", level=1)
    document.add_paragraph(
        "Open Fixture Authors. A Synthetic Record for Reproducible Review Tests. "
        "Public Fixture Laboratory, 2026."
    )
    add_page_number_footer(document)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def paragraph_text(paragraph: etree._Element) -> str:
    chunks: list[str] = []
    for node in paragraph.iter():
        if node.tag in {qnw("t"), qnw("delText")}:
            chunks.append(node.text or "")
    return "".join(chunks)


def find_paragraph(root: etree._Element, contains: str) -> etree._Element:
    matches = [p for p in root.findall(".//w:p", namespaces=NS) if contains in paragraph_text(p)]
    if len(matches) != 1:
        raise RuntimeError(f"expected one paragraph containing {contains!r}; found {len(matches)}")
    return matches[0]


def text_node(text: str, *, deleted: bool = False) -> etree._Element:
    node = etree.Element(qnw("delText" if deleted else "t"))
    if text[:1].isspace() or text[-1:].isspace() or "  " in text:
        node.set(f"{{{XML_NS}}}space", "preserve")
    node.text = text
    return node


def make_run(
    text: str,
    *,
    deleted: bool = False,
    bold: bool = False,
    format_change: tuple[str, str, str] | None = None,
) -> etree._Element:
    run = etree.Element(qnw("r"))
    if bold or format_change is not None:
        rpr = etree.SubElement(run, qnw("rPr"))
        if bold:
            etree.SubElement(rpr, qnw("b"))
        if format_change is not None:
            revision_id, author, timestamp = format_change
            change = etree.SubElement(rpr, qnw("rPrChange"))
            change.set(qnw("id"), revision_id)
            change.set(qnw("author"), author)
            change.set(qnw("date"), timestamp)
            etree.SubElement(change, qnw("rPr"))
    run.append(text_node(text, deleted=deleted))
    return run


def make_revision(
    element_name: str,
    revision_id: str,
    author: str,
    timestamp: str,
    text: str,
    *,
    deleted: bool = False,
) -> etree._Element:
    revision = etree.Element(qnw(element_name))
    revision.set(qnw("id"), revision_id)
    revision.set(qnw("author"), author)
    revision.set(qnw("date"), timestamp)
    revision.append(make_run(text, deleted=deleted))
    return revision


def rewrite_paragraph(paragraph: etree._Element, new_children: list[etree._Element]) -> None:
    ppr = paragraph.find("w:pPr", namespaces=NS)
    starts = list(paragraph.findall("w:bookmarkStart", namespaces=NS))
    ends = list(paragraph.findall("w:bookmarkEnd", namespaces=NS))
    for child in list(paragraph):
        if child is not ppr:
            paragraph.remove(child)
    for node in starts:
        paragraph.append(node)
    for node in new_children:
        paragraph.append(node)
    for node in ends:
        paragraph.append(node)


def split_once(text: str, needle: str) -> tuple[str, str]:
    if text.count(needle) != 1:
        raise RuntimeError(f"expected exactly one {needle!r} in {text!r}")
    return tuple(text.split(needle, 1))  # type: ignore[return-value]


def revision_range_marker(
    name: str,
    range_id: str,
    author: str,
    timestamp: str,
    move_name: str | None = None,
) -> etree._Element:
    marker = etree.Element(qnw(name))
    marker.set(qnw("id"), range_id)
    if name.endswith("RangeStart"):
        marker.set(qnw("author"), author)
        marker.set(qnw("date"), timestamp)
        if move_name is not None:
            marker.set(qnw("name"), move_name)
    return marker


def comment_reference_run(comment_id: str) -> etree._Element:
    run = etree.Element(qnw("r"))
    reference = etree.SubElement(run, qnw("commentReference"))
    reference.set(qnw("id"), comment_id)
    return run


def add_comment_anchor(
    paragraph: etree._Element,
    comment_id: str,
    *,
    point: bool,
) -> None:
    start = etree.Element(qnw("commentRangeStart"))
    start.set(qnw("id"), comment_id)
    end = etree.Element(qnw("commentRangeEnd"))
    end.set(qnw("id"), comment_id)

    children = list(paragraph)
    insert_at = 1 if paragraph.find("w:pPr", namespaces=NS) is not None else 0
    while insert_at < len(children) and children[insert_at].tag == qnw("bookmarkStart"):
        insert_at += 1
    paragraph.insert(insert_at, start)

    if point:
        paragraph.insert(insert_at + 1, end)
        paragraph.insert(insert_at + 2, comment_reference_run(comment_id))
        return

    bookmark_end = paragraph.find("w:bookmarkEnd", namespaces=NS)
    end_at = len(paragraph) if bookmark_end is None else paragraph.index(bookmark_end)
    paragraph.insert(end_at, end)
    paragraph.insert(end_at + 1, comment_reference_run(comment_id))


def apply_known_review_events(root: etree._Element) -> None:

    paragraph = find_paragraph(root, TARGETS["txr_insert_001"])
    prefix, suffix = split_once(paragraph_text(paragraph), "confirms")
    rewrite_paragraph(
        paragraph,
        [
            make_run(prefix),
            make_revision("ins", "101", REVIEWER_ALPHA, TIME_ALPHA, "consistently "),
            make_run("confirms" + suffix),
        ],
    )

    paragraph = find_paragraph(root, TARGETS["txr_delete_001"])
    prefix, suffix = split_once(paragraph_text(paragraph), "redundant ")
    rewrite_paragraph(
        paragraph,
        [
            make_run(prefix),
            make_revision("del", "102", REVIEWER_BETA, TIME_BETA, "redundant ", deleted=True),
            make_run(suffix),
        ],
    )

    paragraph = find_paragraph(root, TARGETS["txr_replace_001"])
    prefix, suffix = split_once(paragraph_text(paragraph), "good")
    rewrite_paragraph(
        paragraph,
        [
            make_run(prefix),
            make_revision("del", "103", REVIEWER_ALPHA, TIME_ALPHA, "good", deleted=True),
            make_revision("ins", "104", REVIEWER_ALPHA, TIME_ALPHA, "consistent"),
            make_run(suffix),
        ],
    )

    paragraph = find_paragraph(root, "Move case: Sensors were calibrated")
    label = "Move case: "
    moved = "The specimen was then loaded. "
    stationary = "Sensors were calibrated before loading. "
    rewrite_paragraph(
        paragraph,
        [
            make_run(label),
            revision_range_marker("moveToRangeStart", "302", REVIEWER_BETA, TIME_BETA, "move-001"),
            make_revision("moveTo", "106", REVIEWER_BETA, TIME_BETA, moved),
            revision_range_marker("moveToRangeEnd", "302", REVIEWER_BETA, TIME_BETA),
            make_run(stationary),
            revision_range_marker(
                "moveFromRangeStart", "301", REVIEWER_BETA, TIME_BETA, "move-001"
            ),
            # ``w:moveFrom`` already supplies the revision semantics. Word
            # requires its moved run content to use ordinary ``w:t``; using
            # ``w:delText`` here passes the Open XML SDK 2.8 schema validator
            # but Microsoft Word reports the package as corrupt.
            make_revision("moveFrom", "105", REVIEWER_BETA, TIME_BETA, moved),
            revision_range_marker("moveFromRangeEnd", "301", REVIEWER_BETA, TIME_BETA),
        ],
    )

    paragraph = find_paragraph(root, TARGETS["txr_format_001"])
    prefix, suffix = split_once(paragraph_text(paragraph), "critical observation")
    rewrite_paragraph(
        paragraph,
        [
            make_run(prefix),
            make_run(
                "critical observation",
                bold=True,
                format_change=("107", REVIEWER_ALPHA, TIME_ALPHA),
            ),
            make_run(suffix),
        ],
    )

    range_para = find_paragraph(root, TARGETS["txr_comment_range_001"])
    add_comment_anchor(range_para, "201", point=False)

    point_para = find_paragraph(root, TARGETS["txr_comment_point_001"])
    add_comment_anchor(point_para, "202", point=True)


def add_comment_record(
    comments_root: etree._Element,
    comment_id: str,
    author: str,
    timestamp: str,
    text: str,
) -> None:
    comment = etree.SubElement(comments_root, qnw("comment"))
    comment.set(qnw("id"), comment_id)
    comment.set(qnw("author"), author)
    comment.set(qnw("date"), timestamp)
    comment_para = etree.SubElement(comment, qnw("p"))
    comment_run = etree.SubElement(comment_para, qnw("r"))
    comment_run.append(text_node(text))


def build_comments_part() -> bytes:
    comments_root = etree.Element(qnw("comments"), nsmap={"w": W_NS})
    add_comment_record(
        comments_root,
        "201",
        REVIEWER_ALPHA,
        TIME_ALPHA,
        "Range comment: verify the normalized residual wording.",
    )
    add_comment_record(
        comments_root,
        "202",
        REVIEWER_BETA,
        TIME_BETA,
        "Point comment: check the sentence boundary.",
    )
    return xml_bytes(comments_root)


def enable_track_revisions(members: dict[str, bytes]) -> None:
    part = "word/settings.xml"
    root = parse_xml_bytes(members[part], part)
    existing = root.findall("w:trackRevisions", namespaces=NS)
    if len(existing) > 1:
        raise ValueError("settings contains duplicate w:trackRevisions elements")
    if not existing:
        track_revisions = etree.Element(qnw("trackRevisions"))
        # CT_Settings requires trackRevisions after proofState and before
        # defaultTabStop. Appending it near ``w:rsids`` passes permissive XML
        # parsers but Microsoft Word rejects the package as unreadable.
        proof_state = root.find("w:proofState", namespaces=NS)
        insert_at = 0 if proof_state is None else root.index(proof_state) + 1
        root.insert(insert_at, track_revisions)
    members[part] = xml_bytes(root)


def add_comments_plumbing(members: dict[str, bytes]) -> None:
    relationships_part = "word/_rels/document.xml.rels"
    relationships = parse_xml_bytes(members[relationships_part], relationships_part)
    existing_relationships = [
        item
        for item in relationships.findall("pr:Relationship", namespaces=NS)
        if item.get("Type") == COMMENTS_REL_TYPE
    ]
    if existing_relationships:
        raise ValueError("base DOCX unexpectedly already contains a comments relationship")

    numeric_ids = []
    for relationship in relationships.findall("pr:Relationship", namespaces=NS):
        relationship_id = relationship.get("Id") or ""
        suffix = relationship_id.removeprefix("rId")
        if suffix.isdigit():
            numeric_ids.append(int(suffix))
    next_relationship_id = f"rId{max(numeric_ids, default=0) + 1}"
    relationship = etree.SubElement(relationships, f"{{{PKG_REL_NS}}}Relationship")
    relationship.set("Id", next_relationship_id)
    relationship.set("Type", COMMENTS_REL_TYPE)
    relationship.set("Target", "comments.xml")
    members[relationships_part] = xml_bytes(relationships)

    content_types_part = "[Content_Types].xml"
    content_types = parse_xml_bytes(members[content_types_part], content_types_part)
    existing_overrides = [
        item
        for item in content_types.findall("ct:Override", namespaces=NS)
        if item.get("PartName") == "/word/comments.xml"
    ]
    if existing_overrides:
        raise ValueError("base DOCX unexpectedly already declares a comments part")
    override = etree.SubElement(content_types, f"{{{CT_NS}}}Override")
    override.set("PartName", "/word/comments.xml")
    override.set("ContentType", COMMENTS_CONTENT_TYPE)
    members[content_types_part] = xml_bytes(content_types)


def build_returned_docx(
    base_docx: Path,
    returned_docx: Path,
) -> None:
    members = read_zip_members(base_docx)
    required_parts = {
        "[Content_Types].xml",
        "word/document.xml",
        "word/settings.xml",
        "word/_rels/document.xml.rels",
    }
    missing_parts = sorted(required_parts - members.keys())
    if missing_parts:
        raise ValueError(f"base DOCX is missing required part(s): {missing_parts}")

    enable_track_revisions(members)
    document_part = "word/document.xml"
    document_root = parse_xml_bytes(members[document_part], document_part)
    apply_known_review_events(document_root)
    members[document_part] = xml_bytes(document_root)
    add_comments_plumbing(members)
    members["word/comments.xml"] = build_comments_part()
    deterministic_zip_members(members, returned_docx)


def write_json(path: Path, data: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(data, handle, ensure_ascii=False, indent=2, sort_keys=True)
        handle.write("\n")


def main() -> int:
    args = parse_args()
    project_root = args.project_root.resolve()
    if not project_root.is_dir():
        raise ValueError(f"project root does not exist: {project_root}")
    fixture_root = (
        args.fixture_root.resolve()
        if args.fixture_root is not None
        else project_root / "tests" / "fixtures" / "e0-minimal-paper"
    )
    fixture_root = assert_within(project_root / "tests" / "fixtures", fixture_root, "fixture root")
    report_root = assert_within(
        project_root / "build", project_root / "build" / "e0-cleanroom", "report root"
    )
    work_root = report_root / "work"
    if work_root.exists():
        shutil.rmtree(work_root)
    work_root.mkdir(parents=True)

    image_path = fixture_root / "source" / "assets" / "response-curve.png"
    raw_base = work_root / "review-base-raw.docx"
    base_docx = fixture_root / "base" / "review-base.docx"
    returned_docx = fixture_root / "returned" / "returned-reviewed.docx"

    generate_response_curve(image_path)
    build_base_docx(raw_base, image_path)
    sanitize_public_docx(raw_base, base_docx)
    build_returned_docx(base_docx, returned_docx)

    generated = {
        str(path.relative_to(project_root).as_posix()): sha256_file(path)
        for path in (image_path, base_docx, returned_docx)
    }
    summary = {
        "schema_version": "e0-cleanroom-build-v1",
        "status": "ok",
        "generator": {
            "script": "scripts/build_e0_public_fixture.py",
            "external_helpers": [],
            "python": ".".join(str(item) for item in sys.version_info[:3]),
            "python_docx": docx.__version__,
            "lxml": etree.__version__,
            "pillow": PIL.__version__,
        },
        "generated": generated,
    }
    write_json(report_root / "build-summary.json", summary)
    print(json.dumps(summary, ensure_ascii=True, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
